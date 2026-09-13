import json
import os
import time
from datetime import datetime

import openpyxl
import requests
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

TOKENS_FILE = "tokens.json"
EXCEL_FILE = "omada_local_users_template.xlsx"
DELETE_SHEET = "EliminarUsuarios"
LOG_SHEET = "Log"


def load_tokens(path=TOKENS_FILE):
    """Carga accessToken, omada_id e interface_access_address generados por omada_connect.py."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"No se encontró {path}. Ejecuta primero omada_connect.py")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def mask(value, visible=4):
    """Enmascara un valor sensible dejando visibles los últimos 'visible' caracteres (visible=0 oculta todo)."""
    if not value:
        return value
    value = str(value)
    if visible <= 0 or len(value) <= visible:
        return "*" * len(value)
    return "*" * (len(value) - visible) + value[-visible:]


def get_sites(base_url, omada_id, access_token, page=1, page_size=100):
    """Obtiene la lista de sites del controlador (GET /openapi/v1/{omadacId}/sites)."""
    url = f"{base_url}/openapi/v1/{omada_id}/sites"
    params = {"page": page, "pageSize": page_size}
    headers = {"Authorization": f"AccessToken={access_token}"}

    resp = requests.get(url, params=params, headers=headers, verify=False)
    resp.raise_for_status()
    data = resp.json()

    if data.get("errorCode") != 0:
        raise RuntimeError(f"Error obteniendo sites: {data}")

    return data["result"]["data"]


def build_site_map(sites):
    """Construye un diccionario {nombre_site: siteId} para resolver site_name -> siteId."""
    return {s["name"]: s["siteId"] for s in sites}


ALL_SITES_KEYWORDS = ("todos", "all", "*")


def resolve_target_sites(site_field, site_map):
    """Resuelve el contenido de la celda site_name a una lista de (nombre_site, siteId).

    Admite tres formatos en la celda:
      - Un solo nombre de site: "SITE_A"
      - Varios nombres separados por coma: "SITE_A, SITE_B"
      - La palabra clave TODOS / ALL / * : aplica a todos los sites del controlador

    Si un nombre de site no existe en el controlador, se incluye igual en el resultado
    con siteId=None, para que el llamador lo reporte como error en el log.
    """
    if site_field is None or str(site_field).strip() == "":
        return []

    text = str(site_field).strip()

    if text.lower() in ALL_SITES_KEYWORDS:
        return list(site_map.items())

    names = [n.strip() for n in text.split(",") if n.strip()]
    return [(name, site_map.get(name)) for name in names]


def get_local_users(base_url, omada_id, site_id, access_token, page=1, page_size=1000):
    """Obtiene la lista de usuarios locales de un site
    (GET /openapi/v1/{omadacId}/sites/{siteId}/hotspot/localusers)."""
    url = f"{base_url}/openapi/v1/{omada_id}/sites/{site_id}/hotspot/localusers"
    params = {"page": page, "pageSize": page_size}
    headers = {"Authorization": f"AccessToken={access_token}"}

    resp = requests.get(url, params=params, headers=headers, verify=False)
    resp.raise_for_status()
    data = resp.json()

    if data.get("errorCode") != 0:
        raise RuntimeError(f"Error obteniendo usuarios locales del site {site_id}: {data}")

    return data["result"]["data"]


def find_user_id(users, user_name):
    """Busca el id del usuario local que coincide con userName (case-sensitive) en la lista de usuarios."""
    for u in users:
        if u.get("userName") == user_name:
            return u.get("id")
    return None


def delete_local_user(base_url, omada_id, site_id, access_token, user_id):
    """Elimina un usuario local (DELETE /openapi/v1/{omadacId}/sites/{siteId}/hotspot/localusers/{id})."""
    url = f"{base_url}/openapi/v1/{omada_id}/sites/{site_id}/hotspot/localusers/{user_id}"
    headers = {"Authorization": f"AccessToken={access_token}"}

    resp = requests.delete(url, headers=headers, verify=False)
    resp.raise_for_status()
    return resp.json()


def read_delete_rows(path=EXCEL_FILE, sheet=DELETE_SHEET):
    """Lee las filas de la hoja EliminarUsuarios del Excel, devolviendo una lista de diccionarios."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]

    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None and row[1] is None:
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def append_log(path, site_name, user_name, resultado, mensaje, sheet=LOG_SHEET):
    """Agrega una fila al final de la hoja Log del Excel, con fecha y hora de la operación."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "omada_delete_local_users.py",
        "eliminar",
        site_name,
        user_name,
        resultado,
        mensaje,
    ])
    wb.save(path)


def main():
    tokens = load_tokens()
    base_url = tokens["interface_access_address"]
    omada_id = tokens["omada_id"]
    access_token = tokens["accessToken"]

    print(f"Usando accessToken: {mask(access_token)}")

    sites = get_sites(base_url, omada_id, access_token)
    site_map = build_site_map(sites)

    rows = read_delete_rows()
    print(f"Filas a procesar en '{DELETE_SHEET}': {len(rows)}\n")

    users_cache = {}  # evita pedir la lista de usuarios repetidas veces por site

    for row in rows:
        site_field = row.get("site_name")
        user_name = row.get("userName")
        targets = resolve_target_sites(site_field, site_map)

        if not targets:
            msg = "La celda site_name está vacía o no se pudo interpretar."
            print(f"[ERROR] {user_name} -> {msg}")
            append_log(EXCEL_FILE, site_field, user_name, "ERROR", msg)
            continue

        for site_name, site_id in targets:
            if not site_id:
                msg = f"El site '{site_name}' no existe en el controlador."
                print(f"[ERROR] {user_name} -> {msg}")
                append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)
                continue

            try:
                if site_id not in users_cache:
                    users_cache[site_id] = get_local_users(base_url, omada_id, site_id, access_token)

                user_id = find_user_id(users_cache[site_id], user_name)

                if not user_id:
                    msg = "No se encontró un usuario local con ese userName en el site."
                    print(f"[ERROR] {user_name} en {site_name} -> {msg}")
                    append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)
                    continue

                result = delete_local_user(base_url, omada_id, site_id, access_token, user_id)
                if result.get("errorCode") == 0:
                    msg = "Usuario eliminado correctamente."
                    print(f"[OK] {user_name} en {site_name} -> {msg}")
                    append_log(EXCEL_FILE, site_name, user_name, "OK", msg)
                else:
                    msg = f"errorCode={result.get('errorCode')} msg={result.get('msg')}"
                    print(f"[ERROR] {user_name} en {site_name} -> {msg}")
                    append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)
            except requests.exceptions.RequestException as e:
                msg = f"Excepción de red/HTTP: {e}"
                print(f"[ERROR] {user_name} en {site_name} -> {msg}")
                append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)

            time.sleep(0.2)  # margen frente al límite de 10 solicitudes/segundo del controlador


if __name__ == "__main__":
    main()
