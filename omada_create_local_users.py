import json
import os
import time
from datetime import datetime, timedelta

import openpyxl
import requests
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

TOKENS_FILE = "tokens.json"
EXCEL_FILE = "omada_local_users_template.xlsx"
CREATE_SHEET = "CrearUsuarios"
LOG_SHEET = "Log"


def load_tokens(path=TOKENS_FILE):
    """Carga accessToken, omada_id e interface_access_address generados por omada_connect.py."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"No se encontró {path}. Ejecuta primero omada_connect.py")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def mask(value, visible=4):
    """Enmascara un valor sensible dejando visibles los últimos 'visible' caracteres (visible=0 oculta todo)."""
    if not value:
        return value
    value = str(value)
    if visible <= 0 or len(value) <= visible:
        return "*" * len(value)
    return "*" * (len(value) - visible) + value[-visible:]


def get_sites(base_url, omada_id, access_token, page=1, page_size=100):
    """Obtiene la lista de sites del controlador (GET /openapi/v1/{omadacId}/sites)."""
    url = f"{base_url}/openapi/v1/{omada_id}/sites"
    params = {"page": page, "pageSize": page_size}
    headers = {"Authorization": f"AccessToken={access_token}"}

    resp = requests.get(url, params=params, headers=headers, verify=False)
    resp.raise_for_status()
    data = resp.json()

    if data.get("errorCode") != 0:
        raise RuntimeError(f"Error obteniendo sites: {data}")

    return data["result"]["data"]


def build_site_map(sites):
    """Construye un diccionario {nombre_site: siteId} para resolver site_name -> siteId."""
    return {s["name"]: s["siteId"] for s in sites}


ALL_SITES_KEYWORDS = ("todos", "all", "*")


def resolve_target_sites(site_field, site_map):
    """Resuelve el contenido de la celda site_name a una lista de (nombre_site, siteId).

    Admite tres formatos en la celda:
      - Un solo nombre de site: "SITE_A"
      - Varios nombres separados por coma: "SITE_A, SITE_B"
      - La palabra clave TODOS / ALL / * : aplica a todos los sites del controlador

    Si un nombre de site no existe en el controlador, se incluye igual en el resultado
    con siteId=None, para que el llamador lo reporte como error en el log.
    """
    if site_field is None or str(site_field).strip() == "":
        return []

    text = str(site_field).strip()

    if text.lower() in ALL_SITES_KEYWORDS:
        return list(site_map.items())

    names = [n.strip() for n in text.split(",") if n.strip()]
    return [(name, site_map.get(name)) for name in names]


def read_create_rows(path=EXCEL_FILE, sheet=CREATE_SHEET):
    """Lee las filas de la hoja CrearUsuarios del Excel, devolviendo una lista de diccionarios."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]

    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None and row[1] is None:
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def to_bool(value, default=False):
    """Convierte valores de celda (TRUE/FALSE, 1/0, texto) a booleano de Python."""
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "si", "sí", "yes")


def to_int(value, default=0):
    """Convierte una celda a entero, usando un valor por defecto si está vacía."""
    if value is None or value == "":
        return default
    return int(value)


def build_payload(row):
    """Arma el cuerpo JSON (CreateLocalUserOpenApiVO) esperado por la API a partir de una fila del Excel."""
    expiration_days = to_int(row.get("expiration_days"), 0)
    expiration_time = 0
    if expiration_days > 0:
        expiration_dt = datetime.now() + timedelta(days=expiration_days)
        expiration_time = int(expiration_dt.timestamp() * 1000)

    portals_raw = row.get("portals") or ""
    portals = [p.strip() for p in str(portals_raw).split(",") if p.strip()]

    apply_to_all = to_bool(row.get("applyToAllPortals"), True)

    rate_limit_mode = to_int(row.get("rateLimitMode"), 0)

    payload = {
        "userName": row.get("userName"),
        "password": row.get("password"),
        "enable": to_bool(row.get("enable"), True),
        "expirationTime": expiration_time,
        "bindingType": to_int(row.get("bindingType"), 0),
        "macAddress": row.get("macAddress") or "",
        "maxUsers": to_int(row.get("maxUsers"), 1),
        "name": row.get("name") or "",
        "phone": row.get("phone") or "",
        "rateLimit": {
            "mode": rate_limit_mode,
            "rateLimitProfileId": row.get("rateLimitProfileId") or "",
            "customRateLimit": {
                "downLimitEnable": to_bool(row.get("downLimitEnable"), False),
                "downLimit": to_int(row.get("downLimit_Kbps"), 0),
                "upLimitEnable": to_bool(row.get("upLimitEnable"), False),
                "upLimit": to_int(row.get("upLimit_Kbps"), 0),
            },
        },
        "trafficLimitEnable": to_bool(row.get("trafficLimitEnable"), False),
        "trafficLimit": to_int(row.get("trafficLimit_MB"), 0),
        "trafficLimitFrequency": to_int(row.get("trafficLimitFrequency"), 0),
        "portals": portals,
        "logout": to_bool(row.get("logout"), True),
        "applyToAllPortals": apply_to_all,
        "dailyLimitEnable": to_bool(row.get("dailyLimitEnable"), False),
        "dailyLimit": {
            "authTimeout": to_int(row.get("authTimeout"), 1),
            "customTimeout": to_int(row.get("customTimeout"), 0),
            "customTimeoutUnit": to_int(row.get("customTimeoutUnit"), 1),
        },
    }
    return payload


def create_local_user(base_url, omada_id, site_id, access_token, payload):
    """Crea un usuario local (POST /openapi/v1/{omadacId}/sites/{siteId}/hotspot/localusers)."""
    url = f"{base_url}/openapi/v1/{omada_id}/sites/{site_id}/hotspot/localusers"
    headers = {
        "Authorization": f"AccessToken={access_token}",
        "Content-Type": "application/json",
    }

    resp = requests.post(url, json=payload, headers=headers, verify=False)
    resp.raise_for_status()
    return resp.json()


def append_log(path, site_name, user_name, resultado, mensaje, sheet=LOG_SHEET):
    """Agrega una fila al final de la hoja Log del Excel, con fecha y hora de la operación."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "omada_create_local_users.py",
        "crear",
        site_name,
        user_name,
        resultado,
        mensaje,
    ])
    wb.save(path)


def main():
    tokens = load_tokens()
    base_url = tokens["interface_access_address"]
    omada_id = tokens["omada_id"]
    access_token = tokens["accessToken"]

    print(f"Usando accessToken: {mask(access_token)}")

    sites = get_sites(base_url, omada_id, access_token)
    site_map = build_site_map(sites)

    rows = read_create_rows()
    print(f"Filas a procesar en '{CREATE_SHEET}': {len(rows)}\n")

    for row in rows:
        site_field = row.get("site_name")
        user_name = row.get("userName")
        targets = resolve_target_sites(site_field, site_map)

        if not targets:
            msg = "La celda site_name está vacía o no se pudo interpretar."
            print(f"[ERROR] {user_name} -> {msg}")
            append_log(EXCEL_FILE, site_field, user_name, "ERROR", msg)
            continue

        payload = build_payload(row)

        for site_name, site_id in targets:
            if not site_id:
                msg = f"El site '{site_name}' no existe en el controlador."
                print(f"[ERROR] {user_name} -> {msg}")
                append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)
                continue

            try:
                result = create_local_user(base_url, omada_id, site_id, access_token, payload)
                if result.get("errorCode") == 0:
                    new_id = result.get("result", {}).get("id", "")
                    msg = f"Usuario creado correctamente (id={new_id})."
                    print(f"[OK] {user_name} en {site_name} -> {msg}")
                    append_log(EXCEL_FILE, site_name, user_name, "OK", msg)
                else:
                    msg = f"errorCode={result.get('errorCode')} msg={result.get('msg')}"
                    print(f"[ERROR] {user_name} en {site_name} -> {msg}")
                    append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)
            except requests.exceptions.RequestException as e:
                msg = f"Excepción de red/HTTP: {e}"
                print(f"[ERROR] {user_name} en {site_name} -> {msg}")
                append_log(EXCEL_FILE, site_name, user_name, "ERROR", msg)

            time.sleep(0.2)  # margen frente al límite de 10 solicitudes/segundo del controlador


if __name__ == "__main__":
    main()
